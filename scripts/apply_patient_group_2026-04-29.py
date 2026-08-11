"""
4.29 시트 환자군 + 감염군 실제 적용
- 이름+병실 매칭으로 patient_group / period_type / infection_strain 업데이트
- 매칭 실패는 스킵
- 트랜잭션 내 처리
"""
import openpyxl
import sqlite3
import sys, io, re
from datetime import datetime

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

EXCEL_PATH = r'C:/Users/총무구매/Desktop/2026년 호남더선요양병원 일일병실현황(자동 복구됨) (1).xlsx'
SHEET_NAME = '4.29'
DB_PATH = r'd:/hospital-supply-app/prisma/hospital-supply.db'

GROUP_MAP = {
    '최고도': 'HIGHEST', '고도': 'HIGH', '중도': 'MEDIUM', '증도': 'MEDIUM',
    '경도': 'LOW', '선택': 'SELECT', '미평가': 'UNRATED',
    '폐렴': 'PNEUMONIA', '패혈증': 'SEPSIS',
}
INFECTION_MAP = {
    'c': 'CRE', 'C': 'CRE',
    'm': 'MR',  'M': 'MR',
    'v': 'VRE', 'V': 'VRE',
}
PERIOD_GROUPS = {'PNEUMONIA', 'SEPSIS'}

BLOCK_STARTS = [0, 6, 12, 18, 24, 30, 36]

def normalize_room(raw):
    if raw is None: return ''
    m = re.search(r'\d+', str(raw))
    return f"{m.group()}호" if m else ''

def main():
    # 엑셀 추출
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[SHEET_NAME]

    excel_rows = []
    current_room = [''] * 7
    for row_idx, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
        for bi, start in enumerate(BLOCK_STARTS):
            if start + 5 >= len(row): continue
            room_cell = row[start]
            no_cell = row[start + 1]
            name_cell = row[start + 2]
            val_cell = row[start + 5]

            if room_cell:
                n = normalize_room(room_cell)
                if n: current_room[bi] = n

            if not name_cell or not isinstance(no_cell, (int, float)):
                continue
            name = str(name_cell).strip()
            if not name: continue

            val = str(val_cell).strip() if val_cell else ''
            if not val: continue

            excel_rows.append({
                'name': name,
                'room': current_room[bi],
                'val': val,
                'excel_row': row_idx,
                'block': bi,
            })

    # DB 입원중 환자
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()
    cur.execute("""
        SELECT id, name, room_no, patient_group, period_type, infection_strain
        FROM patients
        WHERE status='ADMITTED' AND deleted_at IS NULL
    """)
    db = {(r['name'], r['room_no']): dict(r) for r in cur.fetchall()}

    # 분류
    group_updates = []        # patient_group + period_type
    infection_updates = []    # infection_strain
    skipped_no_match = []
    skipped_unknown = []
    duplicates_keys = set()

    # 동명이인 검출 (같은 name+room이 DB에 2명 이상인 경우)
    cur.execute("""
        SELECT name, room_no, COUNT(*) AS c
        FROM patients
        WHERE status='ADMITTED' AND deleted_at IS NULL
        GROUP BY name, room_no HAVING c > 1
    """)
    for r in cur.fetchall():
        duplicates_keys.add((r['name'], r['room_no']))

    for er in excel_rows:
        if not er['room']: continue
        key = (er['name'], er['room'])

        # 환자군 매핑?
        if er['val'] in GROUP_MAP:
            new_group = GROUP_MAP[er['val']]
            if key in duplicates_keys:
                skipped_no_match.append({**er, 'reason': '동명이인'})
                continue
            patient = db.get(key)
            if not patient:
                skipped_no_match.append({**er, 'reason': 'DB매칭없음'})
                continue
            new_period = new_group if new_group in PERIOD_GROUPS else ''
            if patient['patient_group'] == new_group and (patient['period_type'] or '') == new_period:
                continue  # 변경 불필요
            group_updates.append({
                'id': patient['id'],
                'name': er['name'], 'room': er['room'],
                'old_group': patient['patient_group'], 'new_group': new_group,
                'old_period': patient['period_type'] or '', 'new_period': new_period,
            })
        # 감염군 매핑?
        elif er['val'] in INFECTION_MAP:
            new_strain = INFECTION_MAP[er['val']]
            if key in duplicates_keys:
                skipped_no_match.append({**er, 'reason': '동명이인'})
                continue
            patient = db.get(key)
            if not patient:
                skipped_no_match.append({**er, 'reason': 'DB매칭없음'})
                continue
            if (patient['infection_strain'] or '') == new_strain:
                continue
            infection_updates.append({
                'id': patient['id'],
                'name': er['name'], 'room': er['room'],
                'old_strain': patient['infection_strain'] or '', 'new_strain': new_strain,
            })
        else:
            skipped_unknown.append(er)

    print(f"환자군 업데이트: {len(group_updates)}건")
    print(f"감염군 업데이트: {len(infection_updates)}건")
    print(f"스킵 (매칭없음/동명이인): {len(skipped_no_match)}건")
    print(f"인식불가 값: {len(skipped_unknown)}건")

    # 트랜잭션
    print("\n--- 적용 시작 ---")
    try:
        cur.execute("BEGIN")
        for u in group_updates:
            cur.execute("""
                UPDATE patients SET patient_group=?, period_type=?
                WHERE id=?
            """, (u['new_group'], u['new_period'], u['id']))
        for u in infection_updates:
            cur.execute("""
                UPDATE patients SET infection_strain=?
                WHERE id=?
            """, (u['new_strain'], u['id']))
        conn.commit()
        print(f"✓ 환자군 {len(group_updates)}건 + 감염군 {len(infection_updates)}건 = 총 {len(group_updates)+len(infection_updates)}건 적용")
    except Exception as e:
        conn.rollback()
        print(f"✗ 롤백: {e}")
        raise

    # 검증
    print("\n--- 검증 ---")
    sample_ids = [u['id'] for u in group_updates[:5]] + [u['id'] for u in infection_updates[:5]]
    if sample_ids:
        placeholders = ','.join('?' * len(sample_ids))
        cur.execute(f"SELECT id, name, room_no, patient_group, period_type, infection_strain FROM patients WHERE id IN ({placeholders})", sample_ids)
        for r in cur.fetchall():
            print(f"  {r['room_no']} {r['name']}: group={r['patient_group']} period={r['period_type']} strain={r['infection_strain']}")

    if skipped_no_match:
        print(f"\n[스킵 상세] ({len(skipped_no_match)}건)")
        reason_counts = {}
        for s in skipped_no_match:
            reason_counts[s['reason']] = reason_counts.get(s['reason'], 0) + 1
        for k, v in reason_counts.items():
            print(f"  {k}: {v}건")
        for s in skipped_no_match[:10]:
            print(f"  - {s['room']} {s['name']} ({s['val']}) [{s['reason']}]")

    conn.close()
    print(f"\n[완료] {datetime.now().isoformat()}")

if __name__ == '__main__':
    main()
